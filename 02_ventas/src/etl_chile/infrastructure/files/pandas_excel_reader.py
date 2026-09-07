"""Implementacion concreta del puerto SpreadsheetReader usando pandas/openpyxl.

Sustituye a los Connection Managers EXCEL (Provider Microsoft.ACE.OLEDB) y
al componente Excel Source (AccessMode=0, modo tabla/hoja) de los paquetes
SSIS originales.
"""

from __future__ import annotations

import logging
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

logger = logging.getLogger(__name__)


class PandasExcelReader:
    def read_sheet(self, path: Path, sheet_name: str) -> pd.DataFrame:
        sheet = sheet_name.strip("'").rstrip("$")
        try:
            return pd.read_excel(path, sheet_name=sheet, engine="openpyxl")
        except ValueError:
            # Algunos libros de origen (p.ej. exportados por herramientas de
            # terceros) renombran su unica hoja con un nombre autogenerado
            # distinto del declarado en el .dtsx original (que ya no existe
            # como tal). Si el libro tiene una unica hoja, se usa esa por
            # posicion en vez de fallar.
            actual_sheets = load_workbook(path, read_only=True).sheetnames
            if len(actual_sheets) != 1:
                raise
            logger.warning(
                "La hoja '%s' no existe en %s; el libro tiene una unica hoja "
                "('%s') y se usa esa en su lugar",
                sheet,
                path,
                actual_sheets[0],
            )
            return pd.read_excel(path, sheet_name=0, engine="openpyxl")
