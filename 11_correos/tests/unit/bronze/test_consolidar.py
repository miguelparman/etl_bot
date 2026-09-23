"""Tests de bronze/consolidar.py: filtrado de filas de plantilla vacias, manejo de
archivos con hoja/columnas faltantes, deduplicacion, y conversion de fechas
(FechaHora_UTC_Texto texto ISO -> datetime, FechaHora serial Excel -> datetime)."""

from __future__ import annotations

import io

import openpyxl
import pandas as pd

from bronze.consolidar import consolidar, preparar_bandejas, preparar_registro
from bronze.mappings import COLUMNAS_BANDEJAS, COLUMNAS_REGISTRO


def _libro(filas_registro: list[dict], filas_bandejas: list[dict] | None = None, incluir_bandejas: bool = True) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Registro")
    ws.append(COLUMNAS_REGISTRO)
    for fila in filas_registro:
        ws.append([fila.get(c) for c in COLUMNAS_REGISTRO])

    if incluir_bandejas:
        wb_b = wb.create_sheet("Bandejas")
        wb_b.append(COLUMNAS_BANDEJAS)
        for fila in filas_bandejas or []:
            wb_b.append([fila.get(c) for c in COLUMNAS_BANDEJAS])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


_FILA_BASE = {
    "Bandeja": "SM68@movistar.cl",
    "Tipo": "Entrada",
    "FechaHora_UTC_Texto": "2026-09-17T08:59:36.0000000+00:00",
    "FechaHora": 46282.166389,
    "Contacto": "cliente@example.com",
    "Asunto": "Prueba",
    "ID_Mensaje": "MSG-1",
    "ConversationID": "CONV-1",
    "Asesor": "JUAN PEREZ",
    "EsPrimeraEntrada": 1.0,
    "TieneRespuesta": 0,
    "Tiempo_Primera_Respuesta_Horas": None,
    "Estado": None,
    "Tiempo_Respuesta_Horas": None,
}


def test_filtra_filas_de_plantilla_vacias():
    filas = [
        {**_FILA_BASE, "ID_Mensaje": "MSG-1"},
        {},  # fila de plantilla sin usar: Bandeja vacio
        {},
    ]
    contenido = _libro(filas, [{"Correo_Bandeja": "SM68@movistar.cl", "Asesor": "JUAN PEREZ"}])

    registro_df, bandejas_df = consolidar({"archivo.xlsx": contenido})

    assert len(registro_df) == 1
    assert len(bandejas_df) == 1


def test_archivo_sin_hoja_bandejas_se_omite_solo_esa_hoja():
    contenido = _libro([{**_FILA_BASE, "ID_Mensaje": "MSG-1"}], incluir_bandejas=False)

    registro_df, bandejas_df = consolidar({"archivo.xlsx": contenido})

    assert len(registro_df) == 1
    assert bandejas_df.empty


def test_archivo_con_columna_faltante_se_omite():
    # Genera manualmente un libro con una columna de 'Registro' faltante.
    import openpyxl as _openpyxl

    wb = _openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Registro")
    columnas_incompletas = [c for c in COLUMNAS_REGISTRO if c != "Asesor"]
    ws.append(columnas_incompletas)
    ws.append([_FILA_BASE.get(c) for c in columnas_incompletas])
    buffer = io.BytesIO()
    wb.save(buffer)

    registro_df, bandejas_df = consolidar({"incompleto.xlsx": buffer.getvalue()})

    assert registro_df.empty


def test_deduplica_por_id_mensaje_entre_archivos():
    contenido_a = _libro([{**_FILA_BASE, "ID_Mensaje": "MSG-DUP", "Asesor": "A"}])
    contenido_b = _libro([{**_FILA_BASE, "ID_Mensaje": "MSG-DUP", "Asesor": "B"}])

    registro_df, _ = consolidar({"a.xlsx": contenido_a, "b.xlsx": contenido_b})

    assert len(registro_df) == 1


def test_registro_y_bandejas_incluyen_origen_con_nombre_de_archivo():
    contenido_a = _libro([{**_FILA_BASE, "ID_Mensaje": "MSG-A"}], [{"Correo_Bandeja": "A@movistar.cl"}])
    contenido_b = _libro([{**_FILA_BASE, "ID_Mensaje": "MSG-B"}], [{"Correo_Bandeja": "B@movistar.cl"}])

    registro_df, bandejas_df = consolidar({"Registro_A.xlsx": contenido_a, "Registro_B.xlsx": contenido_b})

    assert registro_df.set_index("ID_Mensaje")["ORIGEN"].to_dict() == {
        "MSG-A": "Registro_A.xlsx",
        "MSG-B": "Registro_B.xlsx",
    }
    assert bandejas_df.set_index("Correo_Bandeja")["ORIGEN"].to_dict() == {
        "A@movistar.cl": "Registro_A.xlsx",
        "B@movistar.cl": "Registro_B.xlsx",
    }


def test_preparar_registro_convierte_fechas_y_descarta_invalidas():
    filas = [
        {**_FILA_BASE, "ID_Mensaje": "MSG-1", "FechaHora_UTC_Texto": "2026-09-17T08:59:36.0000000+00:00"},
        {**_FILA_BASE, "ID_Mensaje": "MSG-2", "FechaHora_UTC_Texto": "no-es-una-fecha"},
    ]
    contenido = _libro(filas)
    registro_df, _ = consolidar({"archivo.xlsx": contenido})

    preparado = preparar_registro(registro_df)

    assert len(preparado) == 1
    assert pd.api.types.is_datetime64_any_dtype(preparado["FechaHora_UTC_Texto"])
    assert preparado.iloc[0]["FechaHora_UTC_Texto"] == pd.Timestamp("2026-09-17 08:59:36")
    assert pd.api.types.is_datetime64_any_dtype(preparado["FechaHora"])


def test_preparar_registro_maneja_fechahora_mezclada_entre_archivos():
    # En datos reales, un archivo trae 'FechaHora' como serial de Excel
    # (float) y otro ya como datetime (una celda con formato de fecha
    # aplicado en el origen) -- al consolidar ambos, la columna queda
    # dtype=object con ambos tipos mezclados. Debe convertir cada valor
    # segun corresponda, sin lanzar excepcion.
    import datetime

    contenido_float = _libro([{**_FILA_BASE, "ID_Mensaje": "MSG-FLOAT", "FechaHora": 46282.166389}])
    contenido_datetime = _libro(
        [{**_FILA_BASE, "ID_Mensaje": "MSG-DATETIME", "FechaHora": datetime.datetime(2026, 9, 17, 2, 51, 39)}]
    )

    registro_df, _ = consolidar({"a.xlsx": contenido_float, "b.xlsx": contenido_datetime})
    preparado = preparar_registro(registro_df)

    assert pd.api.types.is_datetime64_any_dtype(preparado["FechaHora"])
    por_id = preparado.set_index("ID_Mensaje")["FechaHora"]
    assert por_id["MSG-DATETIME"] == pd.Timestamp("2026-09-17 02:51:39")
    assert por_id["MSG-FLOAT"].date() == pd.Timestamp("2026-09-17").date()


def test_preparar_bandejas_convierte_fechas():
    contenido = _libro(
        [{**_FILA_BASE, "ID_Mensaje": "MSG-1"}],
        [
            {
                "Correo_Bandeja": "SM68@movistar.cl",
                "Asesor": "JUAN PEREZ",
                "UltimaRevisionEntrada": "2026-09-22T10:19:40.0000000+00:00",
                "UltimaRevisionSalida": "2026-09-22T09:11:51.0000000+00:00",
            }
        ],
    )
    _, bandejas_df = consolidar({"archivo.xlsx": contenido})

    preparado = preparar_bandejas(bandejas_df)

    assert pd.api.types.is_datetime64_any_dtype(preparado["UltimaRevisionEntrada"])
    assert preparado.iloc[0]["UltimaRevisionEntrada"] == pd.Timestamp("2026-09-22 10:19:40")
