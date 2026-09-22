"""Audita, despues de correr main.py, que la copia haya preservado el
FORMATO de cada Excel: mismas hojas 'Registro' y 'Bandejas', mismas columnas
en cada una. No compara el contenido de las filas ni el hash de bytes: los
'Registro_*.xlsx' son controles de bandeja EN VIVO (las macros
'ActualizarControlBandejas_*.bas' de 'Macros_Coordinadores' los reescriben
continuamente durante la jornada), asi que su contenido cambia entre una
descarga y otra aunque la copia haya funcionado perfectamente -- lo unico
que debe mantenerse estable es la estructura.

Uso (desde cualquier directorio):
    python src/correos/verificar_copia.py
"""

from __future__ import annotations

import io
import logging
import sys
from dataclasses import dataclass
from pathlib import Path

import openpyxl

_MODULE_DIR = Path(__file__).resolve().parent  # src/correos: para los imports de abajo
PROJECT_ROOT = _MODULE_DIR.parent.parent  # 11_correos/: para '.env' y 'logs/'
sys.path.insert(0, str(_MODULE_DIR))

from config import SharePointDestinoSettings, SharePointOrigenSettings, cargar_configuracion
from exceptions import CorreosError
from logging_setup import NOMBRE_LOGGER, configurar_logging
from sharepoint_auth import get_graph_token
from sharepoint_client import SharePointClient

logger = logging.getLogger(NOMBRE_LOGGER)

HOJAS_RELEVANTES = ("Registro", "Bandejas")


@dataclass(frozen=True)
class ResultadoVerificacion:
    nombre: str
    ok: bool
    detalle: str | None = None


def _encabezados(contenido: bytes) -> dict[str, list | None]:
    """Nombre de columna (primera fila) de cada hoja relevante, o None si la
    hoja no existe en ese archivo."""
    wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    try:
        resultado: dict[str, list | None] = {}
        for nombre_hoja in HOJAS_RELEVANTES:
            if nombre_hoja not in wb.sheetnames:
                resultado[nombre_hoja] = None
                continue
            ws = wb[nombre_hoja]
            resultado[nombre_hoja] = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        return resultado
    finally:
        wb.close()


def verificar_copia(
    origen: SharePointOrigenSettings,
    cliente_destino: SharePointClient,
    drive_id_destino: str,
    folder_path_destino: str,
) -> list[ResultadoVerificacion]:
    """Para cada .xlsx de la carpeta origen, descarga ambas versiones
    (origen y destino) y compara las columnas de 'Registro' y 'Bandejas'.
    Un archivo que no se pudo leer en algun lado (por ejemplo, no llego a
    copiarse) tambien se reporta como fallido, con el motivo."""
    token_origen = get_graph_token(origen.tenant_id, origen.client_id, origen.client_secret, origen.timeout_ms)
    cliente_origen = SharePointClient(token_origen, origen.timeout_ms)
    site_id_origen = cliente_origen.resolve_site(origen.hostname, origen.site_path)
    drive_id_origen = cliente_origen.resolve_default_drive(site_id_origen)

    nombres = cliente_origen.list_excel_files(drive_id_origen, origen.folder_path)

    resultados: list[ResultadoVerificacion] = []
    for nombre in nombres:
        try:
            origen_bytes = cliente_origen.download_file(drive_id_origen, f"{origen.folder_path}/{nombre}")
        except CorreosError as exc:
            resultados.append(ResultadoVerificacion(nombre, ok=False, detalle=f"no se pudo leer el origen: {exc}"))
            continue
        try:
            destino_bytes = cliente_destino.download_file(drive_id_destino, f"{folder_path_destino}/{nombre}")
        except CorreosError as exc:
            resultados.append(
                ResultadoVerificacion(nombre, ok=False, detalle=f"no existe (o no se pudo leer) en destino: {exc}")
            )
            continue

        enc_origen = _encabezados(origen_bytes)
        enc_destino = _encabezados(destino_bytes)
        hojas_distintas = [h for h in HOJAS_RELEVANTES if enc_origen[h] != enc_destino[h]]

        if hojas_distintas:
            resultados.append(
                ResultadoVerificacion(nombre, ok=False, detalle=f"columnas distintas en hoja(s): {hojas_distintas}")
            )
        else:
            resultados.append(ResultadoVerificacion(nombre, ok=True))

    return resultados


def main() -> int:
    settings = cargar_configuracion(PROJECT_ROOT)
    configurar_logging(settings.log_dir)

    destino = settings.destino
    token_destino = get_graph_token(destino.tenant_id, destino.client_id, destino.client_secret, destino.timeout_ms)
    cliente_destino = SharePointClient(token_destino, destino.timeout_ms)
    site_id_destino = cliente_destino.resolve_site(destino.hostname, destino.site_path)
    drive_id_destino = cliente_destino.resolve_drive(site_id_destino, destino.drive_name)

    resultados = verificar_copia(settings.origen, cliente_destino, drive_id_destino, destino.folder_path)

    for r in resultados:
        if r.ok:
            logger.info("[OK] %s", r.nombre)
        else:
            logger.error("[DIFIERE] %s: %s", r.nombre, r.detalle)

    fallidos = [r for r in resultados if not r.ok]
    logger.info("Verificacion finalizada: %s de %s archivo(s) OK.", len(resultados) - len(fallidos), len(resultados))
    return 1 if fallidos else 0


if __name__ == "__main__":
    sys.exit(main())
