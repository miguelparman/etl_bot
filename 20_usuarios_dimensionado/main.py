from copy import copy as copy_style
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

# ===========================================
# Archivos
# ===========================================
archivo_origen = r"D:\IRISCENE ENGINEERING CORPORATION SLU\Repositorio Bi - Reporting_BPO\BPO_Chile\9. Dimensionamiento_y_Programación\202609\05. Dimensionado y Programación - Llamadas Usuarios Chile del 14 al 20 de setiembre 2026 - PPP.xlsx"

archivo_destino = r"D:\IRISCENE ENGINEERING CORPORATION SLU\BPO - Insumos\Chile\DIMENSIONADO\Dimensionado y Programación.xlsx"

# ===========================================
# Hora de inicio
# ===========================================
hora_inicio = datetime.now()

# ===========================================
# Abrir libros
# ===========================================
wb_origen = load_workbook(archivo_origen, data_only=True, read_only=True)
ws_origen = wb_origen["Curva"]

wb_destino = load_workbook(archivo_destino)
ws_destino = wb_destino["Hoja1"]


def ultima_fila_vacia(ws, columna):
    """
    Retorna la primera fila vacía de una columna (columna como índice numérico).
    """
    fila = ws.max_row

    while fila > 1 and ws.cell(row=fila, column=columna).value is None:
        fila -= 1

    return fila + 1


def copiar_formato(ws, fila_estilo, fila_destino, columna_inicio, columna_fin):
    """
    Copia el formato (fuente, borde, relleno, alineación, protección y
    formato numérico) de una fila de referencia hacia una fila destino.
    """
    for columna in range(columna_inicio, columna_fin + 1):
        celda_origen = ws.cell(row=fila_estilo, column=columna)
        celda_destino = ws.cell(row=fila_destino, column=columna)

        celda_destino.font = copy_style(celda_origen.font)
        celda_destino.border = copy_style(celda_origen.border)
        celda_destino.fill = copy_style(celda_origen.fill)
        celda_destino.alignment = copy_style(celda_origen.alignment)
        celda_destino.protection = copy_style(celda_origen.protection)
        celda_destino.number_format = celda_origen.number_format


def copiar_rango(ws_origen, ws_destino,
                 fila_inicio, fila_fin,
                 col_inicio, col_fin,
                 columna_destino,
                 fila_destino=None):

    if fila_destino is None:
        fila_destino = ultima_fila_vacia(ws_destino, columna_destino)

    columnas = col_fin - col_inicio + 1
    fila_estilo = fila_destino - 1

    filas_origen = ws_origen.iter_rows(
        min_row=fila_inicio, max_row=fila_fin,
        min_col=col_inicio, max_col=col_fin,
        values_only=True
    )

    for i, valores in enumerate(filas_origen):

        copiar_formato(
            ws_destino,
            fila_estilo,
            fila_destino + i,
            columna_destino,
            columna_destino + columnas - 1
        )

        for j, valor in enumerate(valores):

            ws_destino.cell(
                row=fila_destino + i,
                column=columna_destino + j
            ).value = valor

    return fila_destino + (fila_fin - fila_inicio)


def extender_formulas(ws, fila_inicio, fila_fin, columnas):
    """
    Autocompleta las fórmulas de las columnas indicadas (tomando como
    referencia la fila anterior a fila_inicio) hacia todas las filas nuevas,
    ajustando las referencias relativas y copiando también el formato.
    """
    fila_estilo = fila_inicio - 1

    for columna in columnas:
        celda_referencia = ws.cell(row=fila_estilo, column=columna)

        if not isinstance(celda_referencia.value, str) or not celda_referencia.value.startswith("="):
            continue

        formula_origen = celda_referencia.value
        coord_origen = celda_referencia.coordinate

        for fila in range(fila_inicio, fila_fin + 1):
            copiar_formato(ws, fila_estilo, fila, columna, columna)

            celda_destino = ws.cell(row=fila, column=columna)
            celda_destino.value = Translator(
                formula_origen, origin=coord_origen
            ).translate_formula(celda_destino.coordinate)


def agregar_valor_fijo(ws, fila_estilo, fila_inicio, fila_fin, columna, valor, formato_numero):
    """
    Copia el formato de la fila de referencia y escribe un valor fijo
    (con el formato numérico indicado) en cada fila nueva de la columna.
    """
    for fila in range(fila_inicio, fila_fin + 1):
        copiar_formato(ws, fila_estilo, fila, columna, columna)

        celda_destino = ws.cell(row=fila, column=columna)
        celda_destino.value = valor
        celda_destino.number_format = formato_numero


def actualizar_rango_tabla(ws, fila_final):
    """
    Expande el rango (y el autoFilter) de las tablas de la hoja
    que terminen antes de fila_final, para incluir las filas nuevas.
    """
    for tabla in ws.tables.values():
        col_inicio, fila_inicio, col_fin, fila_fin = range_boundaries(tabla.ref)

        if fila_final <= fila_fin:
            continue

        nuevo_ref = (
            f"{get_column_letter(col_inicio)}{fila_inicio}:"
            f"{get_column_letter(col_fin)}{fila_final}"
        )

        tabla.ref = nuevo_ref

        if tabla.autoFilter is not None:
            tabla.autoFilter.ref = nuevo_ref


# =====================================================
# Bloques de filas a copiar (fila_inicio, fila_fin) en "Curva"
# =====================================================
BLOQUES_FILAS = (
    (3, 50),
    (52, 99),
    (101, 148),
    (150, 197),
    (199, 246),
    (248, 295),
    (297, 344),
)

# =====================================================
# Equivalencia de columnas Curva -> Hoja1
# (col_inicio, col_fin, columna_destino)
# =====================================================
MAPEO_COLUMNAS = (
    (2, 6, 1),     # B:F -> A:E
    (8, 9, 6),     # H:I -> F:G
    (12, 12, 8),   # L   -> H
    (16, 16, 10),  # P   -> J
)

fila_inicio_destino = ultima_fila_vacia(ws_destino, 1)
filas_finales = []

for fila_inicio, fila_fin in BLOQUES_FILAS:
    fila_destino_bloque = ultima_fila_vacia(ws_destino, 1)

    for col_inicio, col_fin, columna_destino in MAPEO_COLUMNAS:
        filas_finales.append(copiar_rango(
            ws_origen,
            ws_destino,
            fila_inicio,
            fila_fin,
            col_inicio,
            col_fin,
            columna_destino,
            fila_destino_bloque
        ))

wb_origen.close()


# ===========================================
# Extender fórmulas (K y L)
# ===========================================
fila_final_destino = max(filas_finales)
extender_formulas(ws_destino, fila_inicio_destino, fila_final_destino, (11, 12))

# ===========================================
# Valor fijo columna I (15%)
# ===========================================
agregar_valor_fijo(
    ws_destino,
    fila_inicio_destino - 1,
    fila_inicio_destino,
    fila_final_destino,
    9,
    0.15,
    "0%"
)

# ===========================================
# Actualizar rango de tabla(s)
# ===========================================
actualizar_rango_tabla(ws_destino, fila_final_destino)

# ===========================================
# Guardar
# ===========================================
wb_destino.save(archivo_destino)

# ===========================================
# Hora de fin y tiempo transcurrido
# ===========================================
hora_fin = datetime.now()
segundos_totales = int((hora_fin - hora_inicio).total_seconds())
horas, resto = divmod(segundos_totales, 3600)
minutos, segundos = divmod(resto, 60)

print("Proceso finalizado correctamente.")
print(f"Hora de inicio: {hora_inicio.strftime('%H:%M:%S')}")
print(f"Hora de fin: {hora_fin.strftime('%H:%M:%S')}")
print(f"Tiempo transcurrido: {horas}h {minutos}m {segundos}s")