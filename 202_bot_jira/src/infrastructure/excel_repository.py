"""Adaptador de repositorio: lee/escribe la plantilla Excel de worklogs."""

from typing import List

import pandas as pd
from openpyxl import load_workbook

from ..domain.models import WorklogOutcome, WorklogRow

_TRACKED_COLUMNS = ("Registrado", "Accion", "WorklogID")


class ExcelWorklogRepository:
    def __init__(self, path: str):
        self._path = path
        self._df = None

    def load(self) -> List[WorklogRow]:
        df = pd.read_excel(self._path, dtype={"Ticket": str, "WorklogID": str})
        for columna in _TRACKED_COLUMNS:
            if columna not in df.columns:
                df[columna] = ""
            df[columna] = df[columna].astype(object)
        self._df = df

        filas = []
        for idx, fila in df.iterrows():
            filas.append(
                WorklogRow(
                    index=idx,
                    ticket_raw=fila["Ticket"],
                    fecha=fila["Fecha"],
                    horas=fila["Horas"],
                    comentario=fila.get("Comentario", ""),
                    accion=str(fila.get("Accion", "")),
                    registrado=str(fila.get("Registrado", "")),
                    worklog_id=str(fila.get("WorklogID", "")),
                )
            )
        return filas

    def record_outcome(self, row: WorklogRow, outcome: WorklogOutcome) -> None:
        if outcome.skipped:
            return
        self._df.at[row.index, "Registrado"] = outcome.registrado
        if outcome.accion is not None:
            self._df.at[row.index, "Accion"] = outcome.accion
        if outcome.worklog_id is not None:
            self._df.at[row.index, "WorklogID"] = outcome.worklog_id

    def save(self) -> None:
        """Actualiza solo las columnas rastreadas en el Excel original, preservando
        el formato de celda (fechas, horas, etc.) del resto de columnas."""
        wb = load_workbook(self._path)
        ws = wb.active

        encabezados = {cell.value: cell.column for cell in ws[1] if cell.value}
        for columna in _TRACKED_COLUMNS:
            if columna not in encabezados:
                nueva_col = ws.max_column + 1
                ws.cell(row=1, column=nueva_col, value=columna)
                encabezados[columna] = nueva_col

        for idx, fila in self._df.iterrows():
            excel_row = idx + 2
            for columna in _TRACKED_COLUMNS:
                valor = fila[columna]
                if pd.isna(valor):
                    valor = ""
                ws.cell(row=excel_row, column=encabezados[columna], value=valor)

        wb.save(self._path)
